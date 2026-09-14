from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.schemas.document import LoadedDocument


def load_txt(file_path: Path) -> LoadedDocument:
    """Load a UTF-8 text file into Harbor's common document format."""

    content = file_path.read_text(
        encoding="utf-8"
    )

    return LoadedDocument(
        source=file_path.name,
        file_type="txt",
        content=content,
        metadata={
            "path": str(file_path),
        },
    )


def load_pdf(file_path: Path) -> LoadedDocument:
    """
    Extract PDF text while preserving page boundaries.

    Page markers are intentionally retained because they can later
    help Harbor generate page-aware citations.
    """

    reader = PdfReader(
        str(file_path)
    )

    pages: list[str] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        page_text = page.extract_text() or ""

        if not page_text.strip():
            continue

        pages.append(
            f"[PAGE {page_number}]\n{page_text}"
        )

    content = "\n\n".join(pages)

    return LoadedDocument(
        source=file_path.name,
        file_type="pdf",
        content=content,
        metadata={
            "path": str(file_path),
            "page_count": len(reader.pages),
        },
    )


def load_docx(file_path: Path) -> LoadedDocument:
    """Extract non-empty paragraphs from a Microsoft Word document."""

    document = Document(
        str(file_path)
    )

    paragraphs = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    content = "\n\n".join(paragraphs)

    return LoadedDocument(
        source=file_path.name,
        file_type="docx",
        content=content,
        metadata={
            "path": str(file_path),
            "paragraph_count": len(paragraphs),
        },
    )


def load_document(
    file_path: Path,
) -> LoadedDocument:
    """
    Dispatch loading to the correct extractor based on file extension.

    Keeping this decision in one place prevents the rest of the
    ingestion pipeline from becoming file-type dependent.
    """

    suffix = file_path.suffix.lower()

    if suffix == ".txt":
        return load_txt(file_path)

    if suffix == ".pdf":
        return load_pdf(file_path)

    if suffix == ".docx":
        return load_docx(file_path)

    raise ValueError(
        f"Unsupported file type: {suffix}"
    )